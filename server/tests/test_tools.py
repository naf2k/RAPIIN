"""Tests for document parsers, classifiers and filesystem verification."""
import os
import shutil
import tempfile
from pathlib import Path

import pytest

from rapiin.permissions import sandbox_root
from rapiin.tools import filesystem, verify
from rapiin.tools.analysis import classify_file, parse_and_summarize
from rapiin.tools.parsing import parse_file


@pytest.fixture()
def sandbox_dir():
    """A temp dir inside the sandbox so permission checks pass."""
    root = sandbox_root()
    d = root / ".rapiin_test_parser"
    if d.exists():
        shutil.rmtree(d)
    d.mkdir(parents=True)
    yield d
    shutil.rmtree(d, ignore_errors=True)


def test_parse_txt_and_csv(sandbox_dir):
    txt = sandbox_dir / "catatan.txt"
    txt.write_text("isi catatan")
    result = parse_file(txt)
    assert result["structure"] == "text"
    assert "isi catatan" in result["extracted_text"]

    csv = sandbox_dir / "data.csv"
    csv.write_text("a,b\n1,2\n")
    result = parse_file(csv)
    assert result["structure"] == "csv"


def test_parse_docx(sandbox_dir):
    try:
        import docx
    except Exception:
        pytest.skip("python-docx tidak tersedia")
    doc = docx.Document()
    doc.add_paragraph("Laporan kuartal satu.")
    path = sandbox_dir / "laporan.docx"
    doc.save(str(path))
    result = parse_file(path)
    assert result["structure"] == "docx"
    assert "Laporan kuartal satu." in result["extracted_text"]


def test_parse_xlsx(sandbox_dir):
    try:
        import openpyxl
    except Exception:
        pytest.skip("openpyxl tidak tersedia")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Bulan", "Total"])
    ws.append(["Jan", 100])
    path = sandbox_dir / "rekap.xlsx"
    wb.save(str(path))
    result = parse_file(path)
    assert result["structure"] == "xlsx"
    assert "Jan" in result["extracted_text"]


def test_classify_file(sandbox_dir):
    img = sandbox_dir / "foto.png"
    img.write_bytes(b"\x89PNG\r\n\x1a\n" + b"0" * 20)
    assert classify_file(img)["category"] == "images"
    txt = sandbox_dir / "laporan.txt"
    txt.write_text("x")
    assert classify_file(txt)["category"] == "documents"


def test_parse_unsupported(sandbox_dir):
    weird = sandbox_dir / "file.xyz"
    weird.write_text("x")
    result = parse_file(weird)
    assert result["structure"] == "unsupported"


def test_verify_move_delete(sandbox_dir):
    src_dir = sandbox_dir / "src"
    dest_dir = sandbox_dir / "dest"
    src_dir.mkdir()
    dest_dir.mkdir()
    f = src_dir / "a.txt"
    f.write_text("hello")

    # move
    result = filesystem.move_files(
        None, user_id=1, arguments={"source": str(f), "destination": str(dest_dir)}
    )
    assert result["executed_count"] == 1
    assert result["verified_count"] == 1
    assert result["summary"][0]["status"] == "VERIFIED"
    assert (dest_dir / "a.txt").exists()
    assert not f.exists()

    # delete with verification
    result = filesystem.delete_files(
        None, user_id=1, arguments={"path": str(dest_dir / "a.txt")}
    )
    assert result["executed_count"] == 1
    assert result["verified_count"] == 1
    assert not (dest_dir / "a.txt").exists()


def test_verify_functions_direct():
    assert verify.verify_delete("/nonexistent/path")["verified"] is True
