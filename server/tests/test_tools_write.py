"""Tests for file_mkdir, file_write and file_edit (server-side implementations)."""
import base64
import shutil
from pathlib import Path

import pytest

from rapiin.permissions import sandbox_root
from rapiin.tools import filesystem


@pytest.fixture()
def sandbox_dir():
    """A temp dir inside the sandbox so permission checks pass."""
    root = sandbox_root()
    d = root / ".rapiin_test_write"
    if d.exists():
        shutil.rmtree(d)
    d.mkdir(parents=True)
    yield d
    shutil.rmtree(d, ignore_errors=True)


def _args(path: Path, **kwargs):
    return {"path": str(path), **kwargs}


# --- file_mkdir -----------------------------------------------------------

def test_mkdir_creates_nested_folders(sandbox_dir):
    target = sandbox_dir / "a" / "b"
    result = filesystem.create_directory(None, user_id=1, arguments=_args(target))
    assert result["verified_count"] == 1
    assert result["created"] is True
    assert target.is_dir()


def test_mkdir_existing_folder_is_idempotent(sandbox_dir):
    result = filesystem.create_directory(None, user_id=1, arguments=_args(sandbox_dir))
    assert result["verified_count"] == 1
    assert result["created"] is False


def test_mkdir_refuses_existing_file(sandbox_dir):
    busy = sandbox_dir / "file.txt"
    busy.write_text("x")
    with pytest.raises(FileExistsError):
        filesystem.create_directory(None, user_id=1, arguments=_args(busy))


def test_mkdir_outside_sandbox_blocked(sandbox_dir):
    with pytest.raises(PermissionError):
        filesystem.create_directory(None, user_id=1, arguments={"path": "/definitely-outside-sandbox-xyz"})


# --- file_write -----------------------------------------------------------

def test_write_text_file(sandbox_dir):
    target = sandbox_dir / "catatan.txt"
    result = filesystem.write_file(None, user_id=1, arguments=_args(target, content="halo\n"))
    assert result["verified_count"] == 1
    assert target.read_text() == "halo\n"


def test_write_refuses_overwrite(sandbox_dir):
    target = sandbox_dir / "ada.txt"
    target.write_text("lama")
    with pytest.raises(FileExistsError):
        filesystem.write_file(None, user_id=1, arguments=_args(target, content="baru"))


def test_write_base64_binary(sandbox_dir):
    target = sandbox_dir / "foto.png"
    payload = base64.b64encode(b"\x89PNG\r\n\x1a\n" + b"0" * 100).decode()
    result = filesystem.write_file(None, user_id=1, arguments=_args(target, content=payload, encoding="base64"))
    assert result["verified_count"] == 1
    assert target.stat().st_size == 108


def test_write_rejects_bad_base64(sandbox_dir):
    with pytest.raises(ValueError):
        filesystem.write_file(
            None, user_id=1,
            arguments=_args(sandbox_dir / "x.bin", content="!!!bukan-base64!!!", encoding="base64"),
        )


def test_write_rejects_oversize(sandbox_dir):
    with pytest.raises(ValueError):
        filesystem.write_file(
            None, user_id=1,
            arguments=_args(sandbox_dir / "besar.txt", content="x" * (filesystem.MAX_WRITE_BYTES + 1)),
        )


# --- file_edit ------------------------------------------------------------

def test_edit_replace_text(sandbox_dir):
    target = sandbox_dir / "doc.txt"
    target.write_text("apel jeruk apel\n")
    result = filesystem.edit_file(
        None, user_id=1,
        arguments=_args(target, operation="replace", old="apel", new="mangga", count=1),
    )
    assert result["verified_count"] == 1
    assert target.read_text() == "mangga jeruk apel\n"
    assert (sandbox_dir / "doc.txt.bak").read_text() == "apel jeruk apel\n"


def test_edit_replace_missing_text_fails(sandbox_dir):
    target = sandbox_dir / "doc.txt"
    target.write_text("apel\n")
    with pytest.raises(ValueError):
        filesystem.edit_file(
            None, user_id=1, arguments=_args(target, operation="replace", old="pisang", new="x"),
        )
    assert target.read_text() == "apel\n"


def test_edit_append_and_insert(sandbox_dir):
    target = sandbox_dir / "list.txt"
    target.write_text("satu\n")
    filesystem.edit_file(None, user_id=1, arguments=_args(target, operation="append", text="dua"))
    assert target.read_text() == "satu\ndua"
    filesystem.edit_file(None, user_id=1, arguments=_args(target, operation="insert", line=1, text="nol"))
    assert target.read_text() == "nol\nsatu\ndua"


def test_edit_refuses_pdf(sandbox_dir):
    target = sandbox_dir / "doc.pdf"
    target.write_bytes(b"%PDF-1.4 fake")
    with pytest.raises(ValueError, match="hanya bisa dibaca"):
        filesystem.edit_file(None, user_id=1, arguments=_args(target, operation="append", text="x"))


def test_edit_missing_file_fails(sandbox_dir):
    with pytest.raises(FileNotFoundError):
        filesystem.edit_file(
            None, user_id=1, arguments=_args(sandbox_dir / "tidak-ada.txt", operation="append", text="x"),
        )


def test_edit_xlsx_cell(sandbox_dir):
    import openpyxl

    target = sandbox_dir / "rekap.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "Realisasi"
    workbook["Realisasi"]["B4"] = "lama"
    workbook.save(str(target))
    result = filesystem.edit_file(
        None, user_id=1,
        arguments=_args(target, sheet="Realisasi", cell="B4", value="baru"),
    )
    assert result["verified_count"] == 1
    check = openpyxl.load_workbook(str(target), read_only=True, data_only=True)
    assert check["Realisasi"]["B4"].value == "baru"
    check.close()


def test_edit_xlsx_missing_sheet_fails(sandbox_dir):
    import openpyxl

    target = sandbox_dir / "rekap.xlsx"
    openpyxl.Workbook().save(str(target))
    with pytest.raises(ValueError, match="Sheet tidak ditemukan"):
        filesystem.edit_file(
            None, user_id=1, arguments=_args(target, sheet="TidakAda", cell="A1", value="x"),
        )
